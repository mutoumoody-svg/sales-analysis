"""销售分析 Excel 导出测试。"""

import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent / "backend"))

from app.services.export_service import ExportService


def test_product_detail_is_separated_by_store():
    service = ExportService()
    same_product_two_stores = [
        {
            "store_id": "store-a",
            "store_name": "天猫-慕咖官方旗舰店",
            "platform": "天猫",
            "product_id": "product-1",
            "sku": "SKU-001",
            "product_name": "测试杯",
            "total_qty": 10,
            "total_revenue": 1000,
            "total_cost": 400,
            "gross_profit": 600,
            "gross_margin_pct": 60,
        },
        {
            "store_id": "store-b",
            "store_name": "抖音-慕咖官方旗舰店",
            "platform": "抖音",
            "product_id": "product-1",
            "sku": "SKU-001",
            "product_name": "测试杯",
            "total_qty": 5,
            "total_revenue": 500,
            "total_cost": 200,
            "gross_profit": 300,
            "gross_margin_pct": 60,
        },
    ]

    output = service.export_sales_data({}, [], same_product_two_stores, [])
    workbook = load_workbook(BytesIO(output.read()), data_only=False)

    assert workbook.sheetnames == ["店铺商品明细"]
    sheet = workbook["店铺商品明细"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:3] == ["店铺ID", "店铺名称", "平台"]
    assert sheet.max_row == 3
    assert sheet.cell(2, 2).value == "天猫-慕咖官方旗舰店"
    assert sheet.cell(3, 2).value == "抖音-慕咖官方旗舰店"
    assert sheet.cell(2, headers.index("SKU") + 1).value == "SKU-001"
    assert sheet.cell(3, headers.index("SKU") + 1).value == "SKU-001"
