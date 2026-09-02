"""
core/reporter.py
Excel 报告生成器

将 analyzer.build_all() 的输出字典写入格式化 Excel 文件。
"""

import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment

# ── 颜色 ──────────────────────────────────────────────────────────
C_DARK_BLUE  = "1F3864"  # 全局总计行
C_LIGHT_BLUE = "BDD7EE"  # 小计行
C_HEADER     = "D9E1F2"  # 表头
C_GOLD       = "FFF2CC"  # 毛利汇总亮色行
C_GREEN      = "E2EFDA"  # 正利润
C_RED        = "FFDCE1"  # 负利润/赠品
C_INPUT      = "FFD966"  # 人工输入格（橙黄）

# ── 样式对象 ──────────────────────────────────────────────────────
THIN = Side(border_style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=False)

# 格式码
FMT_MONEY = '#,##0.00'
FMT_INT   = '#,##0'
FMT_PCT   = '0.00%'
FMT_DATE  = 'yyyy-mm-dd'
FMT_DTIME = 'yyyy-mm-dd hh:mm'
FMT_TEXT  = '@'

# ── 每个 Sheet 的列配置 ──────────────────────────────────────────
# (列名, 格式, 宽度)
SHEET_CONFIG: dict[str, list[tuple]] = {
    "订单明细": [
        ("主订单号",  FMT_TEXT,  22),
        ("子订单号",  FMT_TEXT,  22),
        ("日期",      FMT_DATE,  12),
        ("下单时间",  FMT_DTIME, 18),
        ("支付时间",  FMT_DTIME, 18),
        ("发货时间",  FMT_DTIME, 18),
        ("完成时间",  FMT_DTIME, 18),
        ("支付方式",  FMT_TEXT,  10),
        ("订单状态",  FMT_TEXT,  10),
        ("售后状态",  FMT_TEXT,  10),
        ("商品数量",  FMT_INT,    8),
        ("标价总额",  FMT_MONEY, 12),
        ("应付金额",  FMT_MONEY, 12),
        ("商家收入",  FMT_MONEY, 12),
        ("优惠总额",  FMT_MONEY, 12),
        ("平台优惠",  FMT_MONEY, 10),
        ("商家优惠",  FMT_MONEY, 10),
        ("达人优惠",  FMT_MONEY, 10),
        ("商家改价",  FMT_MONEY, 10),
        ("流量类型",  FMT_TEXT,  12),
        ("达人昵称",  FMT_TEXT,  20),
        ("省",        FMT_TEXT,   8),
        ("市",        FMT_TEXT,   8),
        ("区",        FMT_TEXT,   8),
        ("仓库",      FMT_TEXT,  12),
    ],
    "产品明细": [
        ("主订单号",  FMT_TEXT,  22),
        ("日期",      FMT_DATE,  12),
        ("下单时间",  FMT_DTIME, 18),
        ("SKU编码",   FMT_TEXT,  18),
        ("商品名称",  FMT_TEXT,  40),
        ("是否赠品",  FMT_TEXT,   8),
        ("成本状态",  FMT_TEXT,  10),
        ("销量",      FMT_INT,    8),
        ("标价",      FMT_MONEY, 10),
        ("商家收入",  FMT_MONEY, 12),
        ("单位成本",  FMT_MONEY, 10),
        ("总成本",    FMT_MONEY, 12),
        ("毛利润",    FMT_MONEY, 12),
        ("毛利率",    FMT_PCT,   10),
        ("订单状态",  FMT_TEXT,  10),
        ("流量类型",  FMT_TEXT,  12),
        ("达人昵称",  FMT_TEXT,  20),
    ],
    "SKU汇总": [
        ("SKU编码",   FMT_TEXT,  18),
        ("商品名称",  FMT_TEXT,  40),
        ("排名",      FMT_INT,    6),
        ("订单数",    FMT_INT,    8),
        ("总销量",    FMT_INT,   10),
        ("总销售额",  FMT_MONEY, 14),
        ("总成本",    FMT_MONEY, 14),
        ("总毛利润",  FMT_MONEY, 14),
        ("毛利率",    FMT_PCT,   10),
        ("客单价",    FMT_MONEY, 10),
    ],
    "日期趋势": [
        ("日期",      FMT_DATE,  12),
        ("订单数",    FMT_INT,   10),
        ("总销量",    FMT_INT,   10),
        ("总收入",    FMT_MONEY, 14),
        ("总成本",    FMT_MONEY, 14),
        ("总毛利润",  FMT_MONEY, 14),
        ("毛利率",    FMT_PCT,   10),
    ],
    "流量来源": [
        ("流量类型",  FMT_TEXT,  14),
        ("达人昵称",  FMT_TEXT,  24),
        ("订单数",    FMT_INT,   10),
        ("总销量",    FMT_INT,   10),
        ("总销售额",  FMT_MONEY, 14),
        ("总毛利润",  FMT_MONEY, 14),
        ("销售额占比",FMT_PCT,   12),
        ("毛利率",    FMT_PCT,   10),
    ],
    "发货情况": [
        ("订单状态",  FMT_TEXT,  12),
        ("订单数",    FMT_INT,   10),
        ("总收入",    FMT_MONEY, 14),
        ("占比",      FMT_PCT,   10),
        ("主订单号",  FMT_TEXT,  22),
        ("日期",      FMT_DATE,  12),
        ("下单时间",  FMT_DTIME, 18),
        ("商家收入",  FMT_MONEY, 12),
        ("达人昵称",  FMT_TEXT,  20),
        ("省",        FMT_TEXT,   8),
        ("市",        FMT_TEXT,   8),
    ],
    "售后相关销售": [
        ("指标",      FMT_TEXT,  28),
        ("数值",      FMT_TEXT,  20),
        ("主订单号",  FMT_TEXT,  22),
        ("日期",      FMT_DATE,  12),
        ("下单时间",  FMT_DTIME, 18),
        ("订单状态",  FMT_TEXT,  10),
        ("商品数量",  FMT_INT,    8),
        ("退款金额",  FMT_MONEY, 12),
        ("商家收入",  FMT_MONEY, 12),
        ("省",        FMT_TEXT,   8),
        ("市",        FMT_TEXT,   8),
        ("SKU编码",   FMT_TEXT,  18),
        ("商品名称",  FMT_TEXT,  36),
        ("是否赠品",  FMT_TEXT,   8),
        ("销量",      FMT_INT,    8),
        ("单位成本",  FMT_MONEY, 10),
        ("总成本",    FMT_MONEY, 12),
        ("毛利润",    FMT_MONEY, 12),
    ],
    "毛利汇总": [
        ("指标",      FMT_TEXT,  24),
        ("数值",      FMT_TEXT,  20),
        ("排名",      FMT_INT,    6),
        ("SKU编码",   FMT_TEXT,  18),
        ("商品名称",  FMT_TEXT,  40),
        ("销量",      FMT_INT,   10),
        ("销售额",    FMT_MONEY, 14),
        ("毛利润",    FMT_MONEY, 14),
        ("毛利率",    FMT_PCT,   10),
    ],
    # 月度对比表（动态列，按月份列名）——宽度用默认值，列名在运行时生成
    "月度对比": [
        ("指标",  FMT_TEXT,  28),
    ],
}

# 月份前缀正则，用于从 Sheet 名中剥离 "202604_" 前缀
_RE_MONTH_PREFIX = re.compile(r"^\d{6}_")


# ── 样式工具函数 ──────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, color="000000", size=10) -> Font:
    return Font(bold=bold, color=color, size=size)


def _apply(cell, fill=None, font=None, fmt=None, align=CENTER):
    cell.alignment = align
    cell.border = BORDER
    if fill:  cell.fill  = fill
    if font:  cell.font  = font
    if fmt:   cell.number_format = fmt


def _write_header(ws, row_idx: int, cols: list[tuple]):
    """写表头行（蓝灰背景，粗体）。"""
    for ci, (name, _, _) in enumerate(cols, start=1):
        c = ws.cell(row=row_idx, column=ci, value=name)
        _apply(c, fill=_fill(C_HEADER), font=_font(bold=True))


def _write_total(ws, row_idx: int, values: list, cols: list[tuple]):
    """写深蓝全局总计行。"""
    for ci, (val, (_, fmt, _)) in enumerate(zip(values, cols), start=1):
        c = ws.cell(row=row_idx, column=ci, value=val)
        _apply(c, fill=_fill(C_DARK_BLUE), font=_font(bold=True, color="FFFFFF"), fmt=fmt)


def _write_row(ws, row_idx: int, row_data: dict, cols: list[tuple],
               fill=None, bold=False):
    """写普通数据行。"""
    for ci, (name, fmt, _) in enumerate(cols, start=1):
        val = row_data.get(name, "")
        # 处理 NaN / None
        if pd.isna(val) if not isinstance(val, (str, bool)) else False:
            val = ""
        c = ws.cell(row=row_idx, column=ci, value=val)
        _apply(c, fill=fill, font=_font(bold=bold), fmt=fmt)


def _set_col_widths(ws, cols: list[tuple]):
    for ci, (_, _, width) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = width


def _freeze(ws, col_count: int):
    """冻结首行（表头）。"""
    ws.freeze_panes = f"A3"


# ── 主报告生成函数 ────────────────────────────────────────────────

def generate(sheets: dict[str, pd.DataFrame], output_path: str,
             platform: str = "抖音", store: str = "慕咖") -> None:
    """将分析结果写入格式化 Excel 文件。"""
    wb = Workbook()
    wb.remove(wb.active)  # 删除默认空 Sheet

    for sheet_name, df in sheets.items():
        _write_sheet(wb, sheet_name, df, platform, store)

    _add_charts_to_wb(wb, sheets)   # 图表

    wb.save(output_path)
    print(f"\n[完成] 报告已保存：{output_path}")


def _base_name(name: str) -> str:
    """剥离 '202604_' 这类月份前缀，返回纯 Sheet 基础名称。"""
    return _RE_MONTH_PREFIX.sub("", name)


def _write_sheet(wb: Workbook, name: str, df: pd.DataFrame,
                 platform: str, store: str):
    ws = wb.create_sheet(name)
    base = _base_name(name)   # "202604_毛利汇总" → "毛利汇总"

    # ── 月度对比表：特殊处理（动态列） ───────────────────────────
    if base == "月度对比":
        _write_monthly_comparison(ws, df, platform, store)
        return

    cfg = SHEET_CONFIG.get(base) or SHEET_CONFIG.get(name)
    if cfg is None:
        # 未定义格式的 Sheet：直接写原始数据
        _write_raw(ws, df)
        return

    # 过滤 cfg 中实际存在于 df 的列
    present_cols = [(n, f, w) for n, f, w in cfg if n in df.columns]
    if not present_cols:
        _write_raw(ws, df)
        return

    # ── 第1行：标题信息行 ──────────────────────────────────────────
    ws.cell(row=1, column=1, value=f"{platform} · {store} · {name}").font = \
        _font(bold=True, size=11)
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=max(len(present_cols), 1))
    ws.row_dimensions[1].height = 18

    # ── 第2行：表头 ────────────────────────────────────────────────
    _write_header(ws, 2, present_cols)
    ws.row_dimensions[2].height = 16

    # ── 数据行 ────────────────────────────────────────────────────
    for ri, (_, row) in enumerate(df.iterrows(), start=3):
        row_dict = row.to_dict()
        # 特殊着色规则（用 base 匹配，忽略月份前缀）
        fill = _row_fill(base, row_dict)
        _write_row(ws, ri, row_dict, present_cols, fill=fill)

    # ── 总计行（适用于含数值的 Sheet）────────────────────────────
    _maybe_total_row(ws, df, present_cols, base)

    # ── 列宽 & 冻结 ───────────────────────────────────────────────
    _set_col_widths(ws, present_cols)
    ws.freeze_panes = "A3"

    # ── 毛利汇总：物流输入格高亮 + 综合净利润公式 ────────────────
    if base == "毛利汇总":
        _finalize_profit_sheet(ws, df, present_cols)


def _row_fill(sheet_name: str, row: dict) -> PatternFill | None:
    """按 Sheet 和行内容决定背景色。"""
    if sheet_name == "售后相关销售":
        idx = str(row.get("指标", ""))
        if idx.startswith("──"):
            return _fill(C_LIGHT_BLUE)
        if idx in {"售后订单数", "货物成本损耗", "毛利润影响"}:
            return _fill(C_RED)
        if str(row.get("主订单号", "")) or str(row.get("SKU编码", "")):
            return _fill(C_RED)

    if sheet_name == "产品明细":
        if str(row.get("是否赠品", "")) == "赠品":
            return _fill(C_RED)
        if str(row.get("成本状态", "")) == "待录入":
            return _fill(C_GOLD)   # 黄色 = 成本待录入
        profit = row.get("毛利润", 0)
        if isinstance(profit, (int, float)) and not pd.isna(profit):
            return _fill(C_GREEN) if profit > 0 else (
                _fill(C_RED) if profit < 0 else None
            )
    if sheet_name == "毛利汇总":
        idx_val = str(row.get("指标", ""))
        if idx_val.startswith("──"):
            return _fill(C_LIGHT_BLUE)
        if idx_val in {"总销售额(商家收入)", "总毛利润", "整体毛利率", "合计总成本",
                       "推广费用合计"}:
            return _fill(C_GOLD)
        if idx_val in {"  产品成本（正品）", "  赠品/包装成本"}:
            return _fill(C_LIGHT_BLUE)
        if idx_val in {"快递/物流费用",
                        "  平台罚款", "  其他【客服打款】", "  类目服务费",
                        "  淘宝客", "  品销宝", "  品牌新享", "  百亿补贴",
                        "  淘金币", "  先用后付", "  限时红包"}:
            return _fill(C_INPUT)
        if idx_val == "  基础服务费（6%）":
            return _fill(C_LIGHT_BLUE)
        if idx_val in {"平台费用合计", "合计税费"}:
            return _fill(C_GOLD)
        if idx_val == "  增值税销项（销售×9%）":
            return _fill(C_LIGHT_BLUE)   # 销项（要缴的税）
        if idx_val in {"  进项税：产品成本（×9%）",
                       "  进项税：赠品/包材（×13%）",
                       "  进项税：推广费（×6%）"}:
            return _fill(C_GREEN)         # 进项（可抵扣，绿色）
        if idx_val in {"综合净利润", "综合净利润率",
                       "营销后净利润", "净利润率"}:
            return _fill(C_GREEN)
    if sheet_name == "SKU汇总":
        rank = row.get("排名", 999)
        if isinstance(rank, (int, float)) and rank <= 3:
            return _fill(C_GOLD)
    return None


def _maybe_total_row(ws, df: pd.DataFrame, cols: list[tuple], name: str):
    """在数据最后追加深蓝色总计行（仅限数值类 Sheet）。"""
    numeric_sheets = {"订单明细", "产品明细", "SKU汇总", "日期趋势", "流量来源"}
    if name not in numeric_sheets:
        return

    total_row = {}
    for col_name, fmt, _ in cols:
        if col_name in df.columns and fmt in {FMT_MONEY, FMT_INT}:
            try:
                total_row[col_name] = pd.to_numeric(df[col_name], errors="coerce").sum()
            except Exception:
                total_row[col_name] = ""
        elif col_name == list(cols[0])[0]:
            total_row[col_name] = "合计"
        else:
            total_row[col_name] = ""

    # 修正毛利率（不能简单相加）
    if "毛利率" in total_row:
        revenue_col = next(
            (n for n, f, _ in cols if "收入" in n or "销售额" in n), None
        )
        profit_col  = next(
            (n for n, f, _ in cols if "毛利润" in n), None
        )
        if revenue_col and profit_col:
            rev = pd.to_numeric(df[revenue_col], errors="coerce").sum()
            pft = pd.to_numeric(df[profit_col],  errors="coerce").sum()
            total_row["毛利率"] = pft / rev if rev != 0 else 0

    next_row = ws.max_row + 1
    _write_total(ws, next_row, [total_row.get(n, "") for n, _, _ in cols], cols)


def _write_raw(ws, df: pd.DataFrame):
    """无格式配置时的兜底写法。"""
    for ci, col in enumerate(df.columns, start=1):
        ws.cell(row=1, column=ci, value=col).font = _font(bold=True)
    for ri, (_, row) in enumerate(df.iterrows(), start=2):
        for ci, val in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=val)


# ── 图表 ──────────────────────────────────────────────────────────

def _col_idx(sheet_name: str, col_name: str, df_cols) -> int | None:
    """返回列名在实际写出Sheet中的1-based列号（跳过df中不存在的列）。"""
    cfg = SHEET_CONFIG.get(sheet_name, [])
    present = [n for n, _, _ in cfg if n in df_cols]
    try:
        return present.index(col_name) + 1
    except ValueError:
        return None


def _add_charts_to_wb(wb: Workbook, sheets: dict):
    """为日期趋势、SKU汇总、流量来源三个Sheet各添加一张图表。"""
    _chart_daily_trend(wb, sheets)
    _chart_sku_bar(wb, sheets)
    _chart_traffic_bar(wb, sheets)


def _chart_daily_trend(wb: Workbook, sheets: dict):
    """折线图：每日总收入 vs 总毛利润。"""
    name = "日期趋势"
    if name not in wb.sheetnames:
        return
    df = sheets.get(name)
    if df is None or len(df) == 0:
        return

    ws  = wb[name]
    n   = len(df)
    max_row = 2 + n   # row2=header, rows 3..2+n=data

    ci_date    = _col_idx(name, "日期",    df.columns)
    ci_revenue = _col_idx(name, "总收入",  df.columns)
    ci_profit  = _col_idx(name, "总毛利润", df.columns)
    if not all([ci_date, ci_revenue, ci_profit]):
        return

    chart = LineChart()
    chart.title  = "日销售趋势（总收入 vs 总毛利润）"
    chart.style  = 10
    chart.y_axis.title = "金额（元）"
    chart.width  = 26
    chart.height = 14

    chart.add_data(Reference(ws, min_col=ci_revenue, min_row=2, max_row=max_row),
                   titles_from_data=True)
    chart.add_data(Reference(ws, min_col=ci_profit,  min_row=2, max_row=max_row),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=ci_date, min_row=3, max_row=max_row))

    # 深蓝 / 绿色
    chart.series[0].graphicalProperties.line.solidFill = "1F3864"
    chart.series[0].graphicalProperties.line.width     = 25000
    chart.series[1].graphicalProperties.line.solidFill = "70AD47"
    chart.series[1].graphicalProperties.line.width     = 25000

    ws.add_chart(chart, f"A{max_row + 3}")


def _chart_sku_bar(wb: Workbook, sheets: dict):
    """横向柱状图：Top 10 SKU 销售额 vs 毛利润。"""
    name = "SKU汇总"
    if name not in wb.sheetnames:
        return
    df = sheets.get(name)
    if df is None or len(df) == 0:
        return

    ws    = wb[name]
    top_n = min(10, len(df))   # 只取前10行（已按销售额排序）

    ci_label   = _col_idx(name, "商品名称", df.columns)
    ci_revenue = _col_idx(name, "总销售额", df.columns)
    ci_profit  = _col_idx(name, "总毛利润", df.columns)
    if not all([ci_label, ci_revenue, ci_profit]):
        return

    chart = BarChart()
    chart.type   = "bar"      # 横向，名称不会被截断
    chart.title  = "Top 10 SKU 销售额 vs 毛利润"
    chart.style  = 10
    chart.x_axis.title = "金额（元）"
    chart.width  = 26
    chart.height = 16

    chart.add_data(Reference(ws, min_col=ci_revenue, min_row=2, max_row=2 + top_n),
                   titles_from_data=True)
    chart.add_data(Reference(ws, min_col=ci_profit,  min_row=2, max_row=2 + top_n),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=ci_label, min_row=3, max_row=2 + top_n))

    ws.add_chart(chart, f"A{len(df) + 6}")


def _chart_traffic_bar(wb: Workbook, sheets: dict):
    """横向柱状图：各流量来源销售额。"""
    name = "流量来源"
    if name not in wb.sheetnames:
        return
    df = sheets.get(name)
    if df is None or len(df) == 0:
        return

    ws    = wb[name]
    n_rows = min(15, len(df))

    ci_label   = _col_idx(name, "达人昵称", df.columns)
    ci_revenue = _col_idx(name, "总销售额", df.columns)
    ci_profit  = _col_idx(name, "总毛利润", df.columns)
    if not all([ci_label, ci_revenue, ci_profit]):
        return

    chart = BarChart()
    chart.type   = "bar"
    chart.title  = "流量来源销售额分析"
    chart.style  = 10
    chart.x_axis.title = "金额（元）"
    chart.width  = 26
    chart.height = 16

    chart.add_data(Reference(ws, min_col=ci_revenue, min_row=2, max_row=2 + n_rows),
                   titles_from_data=True)
    chart.add_data(Reference(ws, min_col=ci_profit,  min_row=2, max_row=2 + n_rows),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=ci_label, min_row=3, max_row=2 + n_rows))

    ws.add_chart(chart, f"A{n_rows + 6}")


# ── 月度对比表 ───────────────────────────────────────────────────

def _write_monthly_comparison(ws, df: pd.DataFrame, platform: str, store: str):
    """
    月度对比表：完整复用毛利汇总的行结构和着色规则，每月一列。
    行颜色与格式和 毛利汇总 完全一致。
    """
    if df.empty:
        return

    month_cols = [c for c in df.columns if c != "指标"]
    n_total    = 1 + len(month_cols)

    # 哪些指标行用百分比格式
    PCT_ROWS   = {"整体毛利率", "广告费率", "净利润率", "综合净利润率"}
    # 哪些指标行用整数格式
    INT_ROWS   = {"总订单数", "总销量(正品)"}
    # 哪些指标行用两位小数（ROI）
    RATIO_ROWS = {"广告ROI"}

    # ── 第1行：标题 ────────────────────────────────────────────────
    ws.cell(row=1, column=1,
            value=f"{platform} · {store} · 月度对比").font = _font(bold=True, size=11)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_total)
    ws.row_dimensions[1].height = 18

    # ── 第2行：表头 ────────────────────────────────────────────────
    c = ws.cell(row=2, column=1, value="指标")
    _apply(c, fill=_fill(C_HEADER), font=_font(bold=True), align=LEFT)
    for ci, col in enumerate(month_cols, start=2):
        c = ws.cell(row=2, column=ci, value=col)
        _apply(c, fill=_fill(C_HEADER), font=_font(bold=True))
    ws.row_dimensions[2].height = 16

    # ── 数据行 ────────────────────────────────────────────────────
    for ri, (_, row) in enumerate(df.iterrows(), start=3):
        idx_val  = str(row.get("指标", ""))
        row_dict = {"指标": idx_val}   # 供 _row_fill 判断
        row_fill = _row_fill("毛利汇总", row_dict)  # 完全复用毛利汇总颜色规则

        is_bold = idx_val in {
            "总销售额(商家收入)", "总毛利润", "合计总成本",
            "推广费用合计", "合计税费", "平台费用合计",
            "营销后净利润", "综合净利润",
        }

        # 指标列
        c = ws.cell(row=ri, column=1, value=idx_val)
        _apply(c, fill=row_fill, font=_font(bold=is_bold), align=LEFT)

        # 各月数值列
        for ci, col in enumerate(month_cols, start=2):
            val = row.get(col, "")
            # 空字符串留空
            if val == "" or (isinstance(val, float) and val != val):
                val = ""

            # 数字格式选择
            if idx_val in PCT_ROWS:
                fmt = FMT_PCT
            elif idx_val in INT_ROWS:
                fmt = FMT_INT
            elif idx_val in RATIO_ROWS:
                fmt = "0.00"
            elif isinstance(val, (int, float)) and val != "":
                fmt = FMT_MONEY
            else:
                fmt = FMT_TEXT

            c = ws.cell(row=ri, column=ci, value=val if val != "" else None)
            _apply(c, fill=row_fill, font=_font(bold=is_bold), fmt=fmt)

    # ── 列宽 & 冻结 ────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 30
    for ci in range(2, 2 + len(month_cols)):
        ws.column_dimensions[get_column_letter(ci)].width = 18
    ws.freeze_panes = "B3"


# ── 毛利汇总：物流费用后处理 ──────────────────────────────────────

def _finalize_profit_sheet(ws, df: pd.DataFrame, present_cols: list):
    """
    为毛利汇总表后处理（所有单元格已由 _write_row 写好，此处只做覆盖/追加）：
      1. 输入格（物流、各平台费用）：橙黄高亮 + 粗边框 + 批注
      2. 平台费用合计：写 SUM 公式
      3. 综合净利润：= 净利润 - 物流 - 平台费用合计
      4. 综合净利润率：= 综合净利润 / 总销售额
    """
    col_names = [n for n, _, _ in present_cols]
    try:
        b_col = col_names.index("数值") + 1   # "数值"列 1-based
    except ValueError:
        return

    b = get_column_letter(b_col)

    # ── 扫描所有关键行号 ──────────────────────────────────────────
    row_logistics      = None   # 快递/物流费用
    row_total_profit   = None   # 总毛利润
    row_net_profit     = None   # 营销后净利润
    row_net_rate       = None   # 净利润率
    row_composite      = None   # 综合净利润
    row_comp_rate      = None   # 综合净利润率
    row_revenue        = None   # 总销售额（算利润率分母）
    row_platform_items = []     # 各平台费用行号（用于 SUM）
    row_platform_total = None   # 平台费用合计
    row_tax_total      = None   # 合计税费
    row_tax_items      = []     # 税费明细行号（用于 SUM）
    row_ads_items      = []     # 推广场景行号（用于 SUM）
    row_ads_total      = None   # 推广费用合计

    PLATFORM_ITEMS    = {"  平台罚款", "  其他【客服打款】", "  类目服务费"}
    PLATFORM_AUTO     = {"  基础服务费（6%）"}
    TAX_ITEMS         = {
        "  增值税销项（销售×9%）",
        "  进项税：产品成本（×9%）",
        "  进项税：赠品/包材（×13%）",
        "  进项税：推广费（×6%）",
    }
    PROMO_MANUAL_ITEMS = {
        "  淘宝客", "  品销宝", "  品牌新享", "  百亿补贴",
        "  淘金币", "  先用后付", "  限时红包",
    }
    INPUT_ITEMS       = {"快递/物流费用"} | PLATFORM_ITEMS | PROMO_MANUAL_ITEMS

    in_ads_section = False   # 用于捕获推广场景明细行

    for ri, (_, row) in enumerate(df.iterrows(), start=3):
        idx = str(row.get("指标", ""))
        if idx.startswith("── 推广费用"):
            in_ads_section = True
        elif idx == "推广费用合计":
            row_ads_total  = ri
            in_ads_section = False
        elif in_ads_section and idx.startswith("  "):
            row_ads_items.append(ri)   # 各推广场景行（如 "  精准投放"）
        elif idx == "快递/物流费用":
            row_logistics = ri
        elif idx in PLATFORM_ITEMS or idx in PLATFORM_AUTO:
            row_platform_items.append(ri)
        elif idx == "平台费用合计":
            row_platform_total = ri
        elif idx in TAX_ITEMS:
            row_tax_items.append(ri)   # 增值税销项 & 进项抵扣
        elif idx == "合计税费":
            row_tax_total = ri
        elif idx == "总毛利润":
            row_total_profit = ri
        elif idx == "营销后净利润":
            row_net_profit = ri
        elif idx == "净利润率":
            row_net_rate = ri
        elif idx == "综合净利润":
            row_composite = ri
        elif idx == "综合净利润率":
            row_comp_rate = ri
        elif idx == "总销售额(商家收入)":
            row_revenue = ri

    thick = Side(border_style="medium", color="C07000")
    thick_border = Border(left=thick, right=thick, top=thick, bottom=thick)
    thin  = Side(border_style="thin",  color="B0B0B0")
    thin_border  = Border(left=thin,  right=thin,  top=thin,  bottom=thin)

    # ── 1. 各输入格：粗橙边框 + 批注 ────────────────────────────
    hints = {
        "快递/物流费用":      "填入本月快递/物流总费用（元）",
        "  平台罚款":         "填入平台罚款金额（元），无则填 0",
        "  其他【客服打款】": "填入客服打款等其他费用（元）",
        "  类目服务费":       "填入类目服务费（元）",
        "  淘宝客":           "填入本月淘宝客推广费用（元）",
        "  品销宝":           "填入本月品销宝费用（元）",
        "  品牌新享":         "填入本月品牌新享费用（元）",
        "  百亿补贴":         "填入本月百亿补贴费用（元）",
        "  淘金币":           "填入本月淘金币费用（元）",
        "  先用后付":         "填入本月先用后付费用（元）",
        "  限时红包":         "填入本月限时红包费用（元）",
    }
    for ri, (_, row) in enumerate(df.iterrows(), start=3):
        idx = str(row.get("指标", ""))
        if idx in INPUT_ITEMS:
            cell = ws.cell(row=ri, column=b_col)
            cell.fill   = PatternFill("solid", fgColor=C_INPUT)
            cell.border = thick_border
            cell.font   = Font(bold=False, size=10)
            cell.number_format = FMT_MONEY
            if idx in hints:
                note = Comment(hints[idx] + "\n填写后相关合计将自动更新", "系统")
                note.width, note.height = 240, 55
                cell.comment = note

    # ── 2. 平台费用合计 = SUM(各输入格) ──────────────────────────
    if row_platform_total and row_platform_items:
        refs = ",".join(f"{b}{r}" for r in row_platform_items)
        cell = ws.cell(row=row_platform_total, column=b_col)
        cell.value         = f"=SUM({refs})"
        cell.number_format = FMT_MONEY
        cell.font          = Font(bold=True, size=10)
        cell.fill          = PatternFill("solid", fgColor=C_GOLD)
        cell.border        = thin_border

    # ── 2b. 推广费用合计 = SUM(各场景行 + 手工推广项) ────────────
    if row_ads_total and row_ads_items:
        refs = ",".join(f"{b}{r}" for r in row_ads_items)
        cell = ws.cell(row=row_ads_total, column=b_col)
        cell.value         = f"=SUM({refs})"
        cell.number_format = FMT_MONEY
        cell.font          = Font(bold=True, size=10)
        cell.fill          = PatternFill("solid", fgColor=C_GOLD)
        cell.border        = thin_border

    # ── 2d. 营销后净利润 = 总毛利润 - 推广费用合计 ───────────────
    if row_net_profit and row_total_profit and row_ads_total:
        cell = ws.cell(row=row_net_profit, column=b_col)
        cell.value         = f"={b}{row_total_profit}-{b}{row_ads_total}"
        cell.number_format = FMT_MONEY
        cell.font          = Font(bold=True, size=10)
        cell.fill          = PatternFill("solid", fgColor=C_GREEN)
        cell.border        = thin_border

    # ── 2e. 净利润率 = 营销后净利润 / 总销售额 ───────────────────
    if row_net_rate and row_net_profit and row_revenue:
        cell = ws.cell(row=row_net_rate, column=b_col)
        cell.value         = f"=IF({b}{row_revenue}<>0,{b}{row_net_profit}/{b}{row_revenue},0)"
        cell.number_format = FMT_PCT
        cell.font          = Font(bold=True, size=10)
        cell.fill          = PatternFill("solid", fgColor=C_GREEN)
        cell.border        = thin_border

    # ── 2c. 合计税费 = SUM(增值税销项 + 进项抵扣) ───────────────
    if row_tax_total and row_tax_items:
        refs = ",".join(f"{b}{r}" for r in row_tax_items)
        cell = ws.cell(row=row_tax_total, column=b_col)
        cell.value         = f"=SUM({refs})"
        cell.number_format = FMT_MONEY
        cell.font          = Font(bold=True, size=10)
        cell.fill          = PatternFill("solid", fgColor=C_GOLD)
        cell.border        = thin_border

    # ── 3. 综合净利润 = 营销后净利润 - 物流 - 平台费用合计 - 合计税费 ──
    base_row = row_net_profit or row_total_profit   # 优先用营销后净利润
    if row_composite and base_row:
        parts = [f"{b}{base_row}"]
        if row_logistics:      parts.append(f"{b}{row_logistics}")
        if row_platform_total: parts.append(f"{b}{row_platform_total}")
        if row_tax_total:      parts.append(f"{b}{row_tax_total}")
        cell = ws.cell(row=row_composite, column=b_col)
        cell.value         = "=" + "-".join(parts)
        cell.number_format = FMT_MONEY
        cell.font          = Font(bold=True, size=10)
        cell.fill          = PatternFill("solid", fgColor=C_GREEN)
        cell.border        = thin_border

    # ── 4. 综合净利润率 = 综合净利润 / 总销售额 ──────────────────
    if row_comp_rate and row_composite and row_revenue:
        cell = ws.cell(row=row_comp_rate, column=b_col)
        cell.value         = f"=IF({b}{row_revenue}<>0,{b}{row_composite}/{b}{row_revenue},0)"
        cell.number_format = FMT_PCT
        cell.font          = Font(bold=True, size=10)
        cell.fill          = PatternFill("solid", fgColor=C_GREEN)
        cell.border        = thin_border
