"""
Export Service - 数据导出服务.

支持将各分析结果导出为Excel文件。
使用openpyxl生成.xlsx文件。
"""

import io
import json
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _safe_value(v: Any) -> Any:
    """将非基本类型(list/dict/None)转成字符串，避免openpyxl报错."""
    if v is None:
        return ""
    if isinstance(v, (str, int, float, bool)):
        return v
    if isinstance(v, (list, tuple)):
        return json.dumps(v, ensure_ascii=False, default=str)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False, default=str)
    return str(v)


class ExportService:
    """Excel导出服务."""

    # 样式定义
    HEADER_FONT = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    CELL_FONT = Font(name="微软雅黑", size=10)
    NUM_FORMAT = '#,##0.00'
    PCT_FORMAT = '0.0"%"'
    HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    CELL_ALIGN = Alignment(horizontal="left", vertical="center")
    NUM_ALIGN = Alignment(horizontal="right", vertical="center")
    THIN_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def _write_sheet(
        self,
        wb: Workbook,
        sheet_name: str,
        headers: List[str],
        rows: List[Dict[str, Any]],
        col_widths: List[int] = None,
    ):
        """写入一个工作表."""
        ws = wb.create_sheet(title=sheet_name[:31])  # Excel sheet name max 31 chars

        # 写表头
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGN
            cell.border = self.THIN_BORDER

        # 写数据
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, header in enumerate(headers, 1):
                value = _safe_value(row_data.get(header, ""))
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = self.CELL_FONT
                cell.border = self.THIN_BORDER
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.alignment = self.NUM_ALIGN
                    cell.number_format = self.NUM_FORMAT
                else:
                    cell.alignment = self.CELL_ALIGN

        # 设置列宽
        if col_widths:
            for i, width in enumerate(col_widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = width
        else:
            for col_idx, header in enumerate(headers, 1):
                max_len = len(str(header))
                for row_data in rows:
                    val_len = len(str(row_data.get(header, "")))
                    if val_len > max_len:
                        max_len = val_len
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

        # 冻结首行
        ws.freeze_panes = "A2"

        return ws

    def export_abc(self, data: dict) -> io.BytesIO:
        """导出ABC分析."""
        wb = Workbook()
        wb.remove(wb.active)

        # 汇总表
        summary = data.get("summary", {})
        self._write_sheet(
            wb,
            "ABC汇总",
            ["指标", "A类", "B类", "C类", "合计"],
            [{
                "指标": "SKU数量",
                "A类": summary.get("a_count", 0),
                "B类": summary.get("b_count", 0),
                "C类": summary.get("c_count", 0),
                "合计": summary.get("total_skus", 0),
            }, {
                "指标": "销售额",
                "A类": summary.get("a_revenue", 0),
                "B类": summary.get("b_revenue", 0),
                "C类": summary.get("c_revenue", 0),
                "合计": summary.get("total_revenue", 0),
            }, {
                "指标": "利润",
                "A类": summary.get("a_profit", 0),
                "B类": summary.get("b_profit", 0),
                "C类": summary.get("c_profit", 0),
                "合计": summary.get("total_profit", 0),
            }, {
                "指标": "收入占比(%)",
                "A类": summary.get("a_revenue_pct", 0),
                "B类": summary.get("b_revenue_pct", 0),
                "C类": summary.get("c_revenue_pct", 0),
                "合计": 100.0,
            }],
        )

        # 明细表
        items = data.get("items", [])
        if items:
            self._write_sheet(
                wb,
                "ABC明细",
                ["SKU", "产品名称", "品牌", "分类", "等级", "销量", "销售额", "利润", "成本", "毛利率(%)", "累计占比(%)"],
                items,
            )

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def export_gmroi(self, data: dict) -> io.BytesIO:
        """导出GMROI分析."""
        wb = Workbook()
        wb.remove(wb.active)

        summary = data.get("summary", {})
        self._write_sheet(
            wb,
            "GMROI汇总",
            ["指标", "数值"],
            [
                {"指标": "总SKU数", "数值": summary.get("total_skus", 0)},
                {"指标": "总库存成本", "数值": summary.get("total_inv_cost", 0)},
                {"指标": "总利润", "数值": summary.get("total_profit", 0)},
                {"指标": "整体GMROI(%)", "数值": summary.get("overall_gmroi", 0)},
                {"指标": "正回报SKU数", "数值": summary.get("positive_gmroi_count", 0)},
                {"指标": "负回报SKU数", "数值": summary.get("negative_gmroi_count", 0)},
                {"指标": "无库存SKU数", "数值": summary.get("no_inventory_count", 0)},
            ],
        )

        items = data.get("items", [])
        if items:
            self._write_sheet(
                wb,
                "GMROI明细",
                ["SKU", "产品名称", "品牌", "可售库存", "有效库存", "单位成本", "库存成本", "销量", "销售额", "利润", "GMROI(%)", "周转率", "毛利率(%)", "状态"],
                items,
            )

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def export_forecast(self, data: dict) -> io.BytesIO:
        """导出销售预测."""
        wb = Workbook()
        wb.remove(wb.active)

        # 趋势参数
        trend = data.get("trend", {})
        summary = data.get("summary", {})
        self._write_sheet(
            wb,
            "预测参数",
            ["指标", "数值"],
            [
                {"指标": "趋势斜率", "数值": trend.get("slope", 0)},
                {"指标": "截距", "数值": trend.get("intercept", 0)},
                {"指标": "利润斜率", "数值": trend.get("slope_profit", 0)},
                {"指标": "平均环比增长率(%)", "数值": trend.get("avg_growth_rate", 0)},
                {"指标": "R²拟合度", "数值": trend.get("r_squared", 0)},
                {"指标": "下月预测销售额", "数值": summary.get("next_month_revenue", 0)},
                {"指标": "预测置信度", "数值": summary.get("confidence", "low")},
                {"指标": "数据点数", "数值": summary.get("data_points", 0)},
                {"指标": "趋势方向", "数值": summary.get("trend_direction", "flat")},
            ],
        )

        # 历史+预测
        history = data.get("history", [])
        forecast = data.get("forecast", [])
        combined = []
        for h in history:
            combined.append({**h, "类型": "实际"})
        for f in forecast:
            combined.append({**f, "类型": "预测"})

        if combined:
            self._write_sheet(
                wb,
                "销售预测",
                ["期间", "类型", "销售额", "利润", "成本", "销量", "发货量", "退货量"],
                combined,
            )

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def export_cashflow(self, data: dict) -> io.BytesIO:
        """导出现金流预测."""
        wb = Workbook()
        wb.remove(wb.active)

        # 费用明细
        expenses = data.get("expense_breakdown", [])
        if expenses:
            self._write_sheet(
                wb,
                "月均费用",
                ["费用类型", "月均金额", "占比(%)"],
                expenses,
            )

        # 现金流
        history = data.get("history", [])
        forecast = data.get("forecast", [])
        combined = []
        for h in history:
            combined.append({**h, "类型": "实际"})
        for f in forecast:
            combined.append({**f, "类型": "预测"})

        if combined:
            self._write_sheet(
                wb,
                "现金流预测",
                ["期间", "类型", "现金流入", "成本流出", "费用流出", "净现金流", "累计现金流"],
                combined,
            )

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    # ============================================================
    # 字段名 → 中文表头映射
    # ============================================================

    SALES_OVERVIEW_HEADERS = {
        "brand": "品牌",
        "period": "月份",
        "total_revenue": "实际销售额",
        "total_ship_amount": "发货金额",
        "return_amount": "退货金额",
        "return_rate": "退货率(%)",
        "order_count": "订单数",
        "gross_profit": "毛利",
        "gross_margin_pct": "毛利率(%)",
        "total_cost": "实际成本",
        "commission_cost": "佣金成本",
        "total_discount": "总折扣",
        "store_count": "店铺数",
        "product_count": "商品数",
    }

    STORE_HEADERS = {
        "store_id": "店铺ID",
        "store_name": "店铺名称",
        "platform": "平台",
        "channel": "渠道",
        "ship_qty": "发货量",
        "return_qty": "退货量",
        "net_qty": "实际销量",
        "ship_amount": "发货金额",
        "return_amount": "退货金额",
        "return_rate": "退货率(%)",
        "total_revenue": "实际销售额",
        "total_cost": "实际成本",
        "gross_profit": "毛利",
        "gross_margin_pct": "毛利率(%)",
        "commission_cost": "佣金成本",
    }

    PRODUCT_HEADERS = {
        "product_id": "商品ID",
        "sku": "SKU",
        "product_name": "商品名称",
        "category": "分类",
        "brand": "品牌",
        "ship_qty": "发货量",
        "return_qty": "退货量",
        "total_qty": "实际销量",
        "total_revenue": "实际销售额",
        "total_cost": "实际成本",
        "gross_profit": "毛利",
        "gross_margin_pct": "毛利率(%)",
    }

    DAILY_TREND_HEADERS = {
        "date": "日期",
        "order_count": "订单数",
        "revenue": "销售额",
        "gross_profit": "毛利",
        "gross_margin_pct": "毛利率(%)",
        "discount": "折扣",
    }

    PROFIT_SUMMARY_HEADERS = {
        "period": "月份",
        "revenue": "实际销售额",
        "net_cost": "实际成本",
        "gross_profit": "毛利",
        "gross_margin_pct": "毛利率(%)",
        "ship_profit": "发货利润",
        "ship_qty": "发货量",
        "return_qty": "退货量",
        "net_qty": "实际销量",
        "ship_amount": "发货金额",
        "return_amount": "退货金额",
        "return_rate": "退货率(%)",
        "commission_cost": "佣金成本",
        "shipping_fee": "运费收入",
        "shipping_cost": "运费成本",
        "packaging_cost": "包装成本",
        "discount": "折扣",
        "contribution_profit": "贡献利润",
        "net_profit": "净利润",
        "net_margin_pct": "净利率(%)",
        "order_count": "订单数",
    }

    INVENTORY_ITEM_HEADERS = {
        "sku": "SKU",
        "product_name": "商品名称",
        "category": "分类",
        "brand": "品牌",
        "warehouse": "仓库",
        "physical_qty": "物理库存",
        "available_qty": "可用库存",
        "effective_qty": "有效库存",
        "in_transit_qty": "在途库存",
        "reserved_qty": "预占库存",
        "inbound_qty": "入库数量",
        "unit_cost": "单位成本",
        "capital_occupied": "占用资金",
        "monthly_sales": "近月销售",
        "avg_daily_sales": "日均销量",
        "daily_rate": "加权日销",
        "weighted_daily_rate": "加权日均销量",
        "stock_days": "可售天数",
        "days_of_supply": "可供应天数",
        "turnover_days": "周转天数",
        "turnover_category": "周转分类",
        "is_stale": "是否滞销",
        "is_low_stock": "是否低库存",
        "stale_days": "滞销天数",
        "needs_reorder": "需补货",
        "reorder_qty": "建议补货量",
        "reorder_value": "补货金额",
        "priority": "优先级",
        "trend_pct": "趋势百分比",
        "trend_direction": "趋势方向",
        "safety_factor": "安全系数",
        "safety_stock": "安全库存",
        "cycle_demand": "周期需求",
        "last_30d_sales": "近30天销量",
        "last_60d_sales": "近60天销量",
        "last_90d_sales": "近90天销量",
        "total_sold_qty": "累计销量",
        "last_sale_date": "最后销售日期",
        "sales_type": "销售类型",
        "status": "库存状态",
        "risk_level": "风险等级",
    }

    # 利润分析 - 按店铺（字段比销售按店铺多：order_count/revenue/net_cost/shipping_cost/packaging_cost/contribution_profit/net_profit）
    PROFIT_STORE_HEADERS = {
        "store_id": "店铺ID",
        "store_name": "店铺名称",
        "platform": "平台",
        "ship_qty": "发货量",
        "return_qty": "退货量",
        "net_qty": "实际销量",
        "ship_amount": "发货金额",
        "return_amount": "退货金额",
        "return_rate": "退货率(%)",
        "order_count": "订单数",
        "revenue": "实际销售额",
        "net_cost": "实际成本",
        "gross_profit": "毛利",
        "gross_margin_pct": "毛利率(%)",
        "commission_cost": "佣金成本",
        "shipping_cost": "运费成本",
        "packaging_cost": "包装成本",
        "contribution_profit": "贡献利润",
        "net_profit": "净利润",
    }

    # 利润分析 - 按商品（字段：product_id/sku/product_name/category/store_name/ship_qty/return_qty/net_qty/net_revenue/net_cost/ship_profit/net_profit/net_margin_pct）
    PROFIT_PRODUCT_HEADERS = {
        "product_id": "商品ID",
        "sku": "SKU",
        "product_name": "商品名称",
        "category": "分类",
        "store_name": "店铺名称",
        "ship_qty": "发货量",
        "return_qty": "退货量",
        "net_qty": "实际销量",
        "net_revenue": "实际销售额",
        "net_cost": "实际成本",
        "ship_profit": "发货利润",
        "net_profit": "净利润",
        "net_margin_pct": "净利率(%)",
    }

    def _translate_rows(
        self,
        rows: List[Dict[str, Any]],
        header_map: Dict[str, str],
    ) -> tuple:
        """把数据库英文字段名转成中文表头 + 重构数据行.

        返回: (headers: List[str], new_rows: List[Dict])
        """
        if not rows:
            return [], []
        # 按 header_map 顺序排，保留未知字段
        ordered_keys = [k for k in header_map.keys() if k in rows[0]]
        unknown_keys = [k for k in rows[0].keys() if k not in header_map]
        all_keys = ordered_keys + unknown_keys

        headers = [header_map.get(k, k) for k in all_keys]
        new_rows = [{header_map.get(k, k): row.get(k) for k in all_keys} for row in rows]
        return headers, new_rows

    def export_sales_data(
        self,
        overview: dict,
        by_store: list,
        by_product: list,
        daily_trend: list,
    ) -> io.BytesIO:
        """导出销售分析数据."""
        wb = Workbook()
        wb.remove(wb.active)

        # 概览（指标/数值两列）
        if overview:
            overview_rows = []
            for k, v in overview.items():
                if isinstance(v, (dict, list)):
                    continue
                label = self.SALES_OVERVIEW_HEADERS.get(k, k)
                overview_rows.append({"指标": label, "数值": v})
            self._write_sheet(
                wb,
                "销售概览",
                ["指标", "数值"],
                overview_rows,
            )

        if by_store:
            headers, new_rows = self._translate_rows(by_store, self.STORE_HEADERS)
            self._write_sheet(wb, "按店铺", headers, new_rows)

        if by_product:
            headers, new_rows = self._translate_rows(by_product, self.PRODUCT_HEADERS)
            self._write_sheet(wb, "按商品", headers, new_rows)

        if daily_trend:
            headers, new_rows = self._translate_rows(daily_trend, self.DAILY_TREND_HEADERS)
            self._write_sheet(wb, "日趋势", headers, new_rows)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def export_profit_data(
        self,
        summary: dict,
        by_store: list,
        by_product: list,
    ) -> io.BytesIO:
        """导出利润分析数据."""
        wb = Workbook()
        wb.remove(wb.active)

        if summary:
            summary_rows = []
            for k, v in summary.items():
                if isinstance(v, (dict, list)):
                    continue
                label = self.PROFIT_SUMMARY_HEADERS.get(k, k)
                summary_rows.append({"指标": label, "数值": v})
            self._write_sheet(
                wb,
                "利润概览",
                ["指标", "数值"],
                summary_rows,
            )

        if by_store:
            headers, new_rows = self._translate_rows(by_store, self.PROFIT_STORE_HEADERS)
            self._write_sheet(wb, "按店铺利润", headers, new_rows)

        if by_product:
            headers, new_rows = self._translate_rows(by_product, self.PROFIT_PRODUCT_HEADERS)
            self._write_sheet(wb, "按商品利润", headers, new_rows)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def export_inventory_data(
        self,
        summary: dict,
        items: list,
    ) -> io.BytesIO:
        """导出库存分析数据."""
        wb = Workbook()
        wb.remove(wb.active)

        if summary:
            summary_rows = []
            for k, v in summary.items():
                if isinstance(v, (dict, list)):
                    continue
                label = self.INVENTORY_ITEM_HEADERS.get(k, k)
                summary_rows.append({"指标": label, "数值": v})
            self._write_sheet(
                wb,
                "库存概览",
                ["指标", "数值"],
                summary_rows,
            )

        if items:
            headers, new_rows = self._translate_rows(items, self.INVENTORY_ITEM_HEADERS)
            self._write_sheet(wb, "库存明细", headers, new_rows)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf


export_service = ExportService()
